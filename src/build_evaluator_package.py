import csv
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from src.utils import ROOT


FREEZE_DIR = ROOT / "freeze" / "pilot-v1.1"

ROUND1_DATASET = (
    ROOT
    / "analysis"
    / "round1_dataset.csv"
)

MANIFEST = (
    FREEZE_DIR
    / "prepared_manifest.csv"
)

OUTPUT_DIR = (
    ROOT
    / "analysis"
    / "blind_annotation"
)

PRIVATE_DIR = (
    ROOT
    / "analysis"
    / "private"
)

PACKAGE_PATH = (
    OUTPUT_DIR
    / "evaluator_package.xlsx"
)

KEY_PATH = (
    PRIVATE_DIR
    / "blind_key.csv"
)

BLIND_SEED = 20260814


def read_csv(
    path: Path,
) -> list[dict[str, str]]:
    """Read a CSV file."""

    with path.open(
        "r",
        encoding="utf-8",
        newline="",
    ) as file:
        return list(
            csv.DictReader(file)
        )


def read_prompt(
    prompt_file: str,
) -> str:
    """
    Read the user-visible material from a frozen
    materialized Round 1 prompt.
    """

    filename = Path(
        prompt_file
    ).name

    path = (
        FREEZE_DIR
        / "prompts"
        / filename
    )

    if not path.exists():
        raise RuntimeError(
            f"Frozen prompt not found: {path}"
        )

    text = path.read_text(
        encoding="utf-8"
    )

    system_marker = (
        "===== SYSTEM PROMPT ====="
    )

    user_marker = (
        "===== USER PROMPT ====="
    )

    if (
        system_marker not in text
        or user_marker not in text
    ):
        raise RuntimeError(
            f"Invalid prompt structure: {path}"
        )

    _, remainder = text.split(
        system_marker,
        1,
    )

    _, user_text = remainder.split(
        user_marker,
        1,
    )

    return user_text.strip()


def write_private_key(
    rows: list[dict[str, str]],
) -> None:
    """Write the private blind-ID mapping."""

    if not rows:
        raise RuntimeError(
            "Cannot write an empty blind key."
        )

    PRIVATE_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with KEY_PATH.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=list(
                rows[0].keys()
            ),
        )

        writer.writeheader()
        writer.writerows(
            rows
        )


def build_instructions_sheet(
    workbook: Workbook,
) -> None:
    """Create evaluator instructions."""

    sheet = workbook.active
    sheet.title = "Instructions"

    content = [
        (
            "Blind Evidential Grounding Evaluation",
            True,
        ),
        (
            "Purpose",
            True,
        ),
        (
            "Evaluate how well each generated "
            "assessment is grounded in the "
            "information provided.",
            False,
        ),
        (
            "Do not evaluate whether the "
            "trustworthiness attribution is "
            "correct, desirable, morally "
            "appropriate, or consistent with "
            "an expected experimental outcome.",
            False,
        ),
        (
            "The items have been randomized and "
            "blinded. Do not attempt to infer "
            "the model, experimental condition, "
            "repetition, or experimental role "
            "of the case.",
            False,
        ),
        (
            "Variables",
            True,
        ),
        (
            "Grounding score: assign one integer "
            "from 0 to 4. Use the Rubric sheet "
            "for the operational definitions.",
            False,
        ),
        (
            "Limitation recognition: assign 1 "
            "if the assessment explicitly "
            "recognizes at least one material "
            "evidential limitation relevant to "
            "its reasoning; otherwise assign 0.",
            False,
        ),
        (
            "A value of 0 for limitation "
            "recognition does not by itself "
            "imply that the assessment is "
            "defective. The relevance of "
            "limitations depends on the "
            "available material.",
            False,
        ),
        (
            "Unsupported evidence: assign 1 if "
            "the assessment introduces at least "
            "one material factual claim used in "
            "its reasoning that is not supported "
            "by the available material; "
            "otherwise assign 0.",
            False,
        ),
        (
            "Do not mark harmless paraphrases or "
            "reasonable linguistic reformulations "
            "as unsupported evidence. The claim "
            "must be materially factual and "
            "unsupported by the provided material.",
            False,
        ),
        (
            "Annotator notes are optional and "
            "may be used to document difficult, "
            "ambiguous, or borderline decisions.",
            False,
        ),
    ]

    for row_index, (
        text,
        bold,
    ) in enumerate(
        content,
        start=1,
    ):
        cell = sheet.cell(
            row=row_index,
            column=1,
            value=text,
        )

        cell.font = Font(
            bold=bold
        )

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    sheet.column_dimensions[
        "A"
    ].width = 110


def build_rubric_sheet(
    workbook: Workbook,
) -> None:
    """Create the grounding rubric sheet."""

    sheet = workbook.create_sheet(
        "Rubric"
    )

    headers = [
        "Score",
        "Label",
        "Operational description",
    ]

    for column, value in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=1,
            column=column,
            value=value,
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    rubric = [
        (
            0,
            "Ungrounded",
            "The assessment is not meaningfully "
            "grounded in the available information. "
            "It may introduce material unsupported "
            "claims, contradict the evidence, assert "
            "the attribution without connecting it "
            "to the evidence, or fail to provide an "
            "interpretable evidential basis.",
        ),
        (
            1,
            "Superficially grounded",
            "The assessment refers to information "
            "from the case but does not adequately "
            "explain how that information supports "
            "the attribution. Relevant evidence or "
            "material limitations may be ignored, "
            "or the reasoning may remain generic.",
        ),
        (
            2,
            "Partially grounded",
            "The assessment meaningfully relates "
            "relevant evidence to the attribution "
            "but contains important grounding "
            "limitations. It may omit important "
            "limitations, overextend the scope of "
            "the evidence, or handle incomplete or "
            "conflicting evidence only partially.",
        ),
        (
            3,
            "Well grounded",
            "The assessment clearly and "
            "appropriately relates materially "
            "relevant evidence to the attribution, "
            "recognizes relevant evidential "
            "limitations, avoids unsupported "
            "factual information, and keeps the "
            "strength and scope of the conclusion "
            "compatible with the available "
            "evidential basis.",
        ),
        (
            4,
            "Critically grounded",
            "The assessment satisfies the "
            "requirements for score 3 and "
            "additionally demonstrates explicit "
            "critical handling of uncertainty, "
            "conflicting evidence, or important "
            "informational limitations. It explains "
            "how those limitations affect the "
            "strength or scope of the attribution "
            "and, where relevant, why a stronger "
            "or weaker conclusion would not be "
            "justified.",
        ),
    ]

    for row_index, values in enumerate(
        rubric,
        start=2,
    ):
        for column, value in enumerate(
            values,
            start=1,
        ):
            cell = sheet.cell(
                row=row_index,
                column=column,
                value=value,
            )

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    sheet.column_dimensions[
        "A"
    ].width = 10

    sheet.column_dimensions[
        "B"
    ].width = 25

    sheet.column_dimensions[
        "C"
    ].width = 100

    sheet.freeze_panes = "A2"


def build_evaluation_sheet(
    workbook: Workbook,
    rows: list[dict[str, str]],
    prompt_files: dict[str, str],
) -> list[dict[str, str]]:
    """
    Create the blinded evaluation sheet and
    return the private blind-ID mapping.
    """

    sheet = workbook.create_sheet(
        "Evaluation"
    )

    headers = [
        "blind_id",
        "case_material",
        "generated_assessment",
        "grounding_score",
        "limitation_recognition",
        "unsupported_evidence",
        "annotator_notes",
    ]

    for column, value in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=1,
            column=column,
            value=value,
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    randomizer = random.Random(
        BLIND_SEED
    )

    shuffled = list(
        rows
    )

    randomizer.shuffle(
        shuffled
    )

    key_rows = []

    for index, row in enumerate(
        shuffled,
        start=1,
    ):
        blind_id = (
            f"ITEM_{index:03d}"
        )

        run_id = row[
            "run_id"
        ]

        if run_id not in prompt_files:
            raise RuntimeError(
                f"Missing prompt mapping "
                f"for {run_id}."
            )

        case_material = read_prompt(
            prompt_files[
                run_id
            ]
        )

        assessment_path = (
            FREEZE_DIR
            / "parsed"
            / f"{run_id}.json"
        )

        if not assessment_path.exists():
            raise RuntimeError(
                f"Frozen assessment not found: "
                f"{assessment_path}"
            )

        assessment = json.loads(
            assessment_path.read_text(
                encoding="utf-8"
            )
        )

        assessment_text = json.dumps(
            assessment,
            ensure_ascii=False,
            indent=2,
        )

        excel_row = (
            index + 1
        )

        values = [
            blind_id,
            case_material,
            assessment_text,
            None,
            None,
            None,
            None,
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            cell = sheet.cell(
                row=excel_row,
                column=column,
                value=value,
            )

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

        key_rows.append(
            {
                "blind_id": blind_id,
                "run_id": run_id,
                "case_id": row[
                    "case_id"
                ],
                "domain": row[
                    "domain"
                ],
                "condition": row[
                    "condition"
                ],
                "model_slot": row[
                    "model_slot"
                ],
                "repetition": row[
                    "repetition"
                ],
            }
        )

    grounding_validation = (
        DataValidation(
            type="list",
            formula1='"0,1,2,3,4"',
            allow_blank=True,
        )
    )

    binary_validation = (
        DataValidation(
            type="list",
            formula1='"0,1"',
            allow_blank=True,
        )
    )

    sheet.add_data_validation(
        grounding_validation
    )

    sheet.add_data_validation(
        binary_validation
    )

    grounding_validation.add(
        f"D2:D{len(rows) + 1}"
    )

    binary_validation.add(
        f"E2:F{len(rows) + 1}"
    )

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        f"A1:G{len(rows) + 1}"
    )

    sheet.column_dimensions[
        "A"
    ].width = 14

    sheet.column_dimensions[
        "B"
    ].width = 85

    sheet.column_dimensions[
        "C"
    ].width = 85

    sheet.column_dimensions[
        "D"
    ].width = 18

    sheet.column_dimensions[
        "E"
    ].width = 24

    sheet.column_dimensions[
        "F"
    ].width = 22

    sheet.column_dimensions[
        "G"
    ].width = 55

    return key_rows


def main() -> None:
    rows = read_csv(
        ROUND1_DATASET
    )

    if len(rows) != 72:
        raise RuntimeError(
            f"Expected 72 Round 1 rows, "
            f"found {len(rows)}."
        )

    run_ids = {
        row["run_id"]
        for row in rows
    }

    if len(run_ids) != 72:
        raise RuntimeError(
            "Expected 72 unique Round 1 run IDs."
        )

    manifest_rows = read_csv(
        MANIFEST
    )

    if len(manifest_rows) != 72:
        raise RuntimeError(
            f"Expected 72 manifest rows, "
            f"found {len(manifest_rows)}."
        )

    prompt_files = {
        row["run_id"]: row["prompt_file"]
        for row in manifest_rows
    }

    if len(prompt_files) != 72:
        raise RuntimeError(
            "Expected 72 unique prompt mappings."
        )

    if set(
        prompt_files
    ) != run_ids:
        raise RuntimeError(
            "Round 1 dataset and frozen "
            "manifest contain different run IDs."
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    PRIVATE_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    build_instructions_sheet(
        workbook
    )

    build_rubric_sheet(
        workbook
    )

    key_rows = build_evaluation_sheet(
        workbook,
        rows,
        prompt_files,
    )

    workbook.save(
        PACKAGE_PATH
    )

    write_private_key(
        key_rows
    )

    print(
        "Evaluation items:",
        len(rows),
    )

    print(
        "Evaluator package:",
        PACKAGE_PATH,
    )

    print(
        "Private blind key:",
        KEY_PATH,
    )

    print(
        "Blind seed:",
        BLIND_SEED,
    )


if __name__ == "__main__":
    main()
